# doc_search_app.py
# pyinstaller --onefile --windowed doc_search_app.py
import sys
import os
import pickle
import subprocess
from pathlib import Path
from collections import defaultdict
import regex
import docx, pptx, openpyxl
import PyPDF2
from PyQt5.QtWidgets import (
    QApplication, QWidget, QPushButton, QFileDialog,
    QTextEdit, QVBoxLayout, QLineEdit, QLabel, QProgressDialog, QListWidget, QListWidgetItem
)
from PyQt5.QtCore import Qt, QThread, pyqtSignal

INDEX_FILE = 'doc_index.pkl'
LINE_MAP_FILE = 'line_map.pkl'

# ---------------------
# Text Extraction Logic
# ---------------------
def extract_text(path):
    try:
        suffix = Path(path).suffix.lower()
        if suffix == '.docx':
            doc = docx.Document(path)
            return '\n'.join(p.text for p in doc.paragraphs)

        elif suffix == '.pptx':
            prs = pptx.Presentation(path)
            text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)
            return '\n'.join(text)

        elif suffix == '.xlsx':
            wb = openpyxl.load_workbook(path, data_only=True)
            text = []
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    row_text = [str(cell) for cell in row if cell]
                    text.append(' '.join(row_text))
            return '\n'.join(text)

        else:
            return ''
    except Exception as e:
        return ''

# ---------------------
# File Finder
# ---------------------
def find_office_files(directory):
    exts = ['*.pptx', '*.docx', '*.xlsx', '*.pdf']
    results = []
    for ext in exts:
        results.extend(Path(directory).rglob(ext))
    return results

# ---------------------
# Indexing Worker Thread
# ---------------------
class IndexWorker(QThread):
    progress = pyqtSignal(int)
    finished = pyqtSignal(dict, dict, int)

    def __init__(self, files):
        super().__init__()
        self.files = files

    def run(self):
        index = defaultdict(list)
        line_map = defaultdict(dict)
        total = len(self.files)
        for idx, file in enumerate(self.files):
            suffix = Path(file).suffix.lower()
            title = Path(file).stem.lower()
            for word in regex.findall(r'\p{L}+', title):
                index[word.lower()].append((file, -1))

            if suffix != '.pdf':
                content = extract_text(file)
                lines = content.split('\n')
                for i, line in enumerate(lines):
                    for word in regex.findall(r'\p{L}+', line):
                        index[word.lower()].append((file, i))
                    line_map[file][i] = line

            self.progress.emit(int((idx + 1) / total * 100))
        self.finished.emit(index, line_map, total)

# ---------------------
# GUI Class
# ---------------------
class FileFinderApp(QWidget):
    def __init__(self):
        super().__init__()
        self.index = {}
        self.line_map = {}
        self.initUI()
        self.load_index_from_file()

    def initUI(self):
        self.setWindowTitle('문서 검색 앱')
        self.setGeometry(100, 100, 800, 600)

        layout = QVBoxLayout()

        self.searchBar = QLineEdit(self)
        self.searchBar.setPlaceholderText("검색어 입력")
        layout.addWidget(QLabel("검색어:"))
        layout.addWidget(self.searchBar)

        self.searchBtn = QPushButton('검색 실행')
        self.searchBtn.clicked.connect(self.search)
        layout.addWidget(self.searchBtn)

        self.resultList = QListWidget()
        self.resultList.itemDoubleClicked.connect(self.open_file)
        layout.addWidget(self.resultList)

        self.folderBtn = QPushButton('폴더 선택 및 색인')
        self.folderBtn.clicked.connect(self.browse_folder)
        layout.addWidget(self.folderBtn)

        self.setLayout(layout)

    def browse_folder(self):
        folder = QFileDialog.getExistingDirectory(self, '폴더 선택')
        if folder:
            files = find_office_files(folder)

            self.progress = QProgressDialog("색인 생성 중...", None, 0, 100, self)
            self.progress.setWindowTitle("로딩 중")
            self.progress.setWindowModality(Qt.ApplicationModal)
            self.progress.setMinimumDuration(0)
            self.progress.show()

            self.worker = IndexWorker(files)
            self.worker.progress.connect(self.progress.setValue)
            self.worker.finished.connect(self.indexing_done)
            self.worker.start()

    def indexing_done(self, index, line_map, count):
        self.index = index
        self.line_map = line_map
        self.save_index_to_file()
        self.resultList.clear()
        self.resultList.addItem(f"색인 완료 및 저장: {count}개 파일")
        self.progress.close()

    def search(self):
        keyword = self.searchBar.text().strip().lower()
        if not keyword:
            self.resultList.clear()
            self.resultList.addItem("검색어를 입력하세요.")
            return

        results = self.index.get(keyword, [])
        self.resultList.clear()
        if not results:
            self.resultList.addItem("결과 없음")
            return

        for file, line_num in results:
            if line_num == -1:
                display = f"[제목 일치] {file}"
            else:
                line = self.line_map[file].get(line_num, '')
                display = f"{file} (줄 {line_num+1}): {line}"
            item = QListWidgetItem(display)
            item.setData(Qt.UserRole, str(file))
            self.resultList.addItem(item)

    def open_file(self, item):
        path = item.data(Qt.UserRole)
        if os.path.exists(path):
            try:
                if sys.platform == "win32":
                    os.startfile(path)
                elif sys.platform == "darwin":
                    subprocess.call(["open", path])
                else:
                    subprocess.call(["xdg-open", path])
            except Exception as e:
                print(f"파일 열기 실패: {e}")

    def save_index_to_file(self):
        with open(INDEX_FILE, 'wb') as f:
            pickle.dump(self.index, f)
        with open(LINE_MAP_FILE, 'wb') as f:
            pickle.dump(self.line_map, f)

    def load_index_from_file(self):
        if os.path.exists(INDEX_FILE) and os.path.exists(LINE_MAP_FILE):
            with open(INDEX_FILE, 'rb') as f:
                self.index = pickle.load(f)
            with open(LINE_MAP_FILE, 'rb') as f:
                self.line_map = pickle.load(f)

if __name__ == '__main__':
    app = QApplication(sys.argv)
    ex = FileFinderApp()
    ex.show()
    sys.exit(app.exec_())
